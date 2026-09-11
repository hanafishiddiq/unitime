"""
Semantic Text and Narrative Slicer for UniTime AI Ingestion Gateway.

Splits long narrative memos, syllabi, curriculum documents, and administrative policies
into coherent semantic chunks along headers (#, ##), department/faculty boundaries,
and double line breaks.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional, Pattern, Set, Tuple, Union

try:
    import docx  # python-docx for .docx support
except ImportError:
    docx = None  # type: ignore[assignment]

from .models import TextChunk


class TextSlicer:
    """
    Slices narrative curriculum documents, academic memos, and policy guidelines
    into semantically coherent text chunks.

    Supports:
    - Markdown headers (#, ##, ###)
    - Department, Faculty, Program Studi (Prodi) administrative boundaries
    - Legal/policy sections (BAB, PASAL, LAMPIRAN, MEMO, SUBJECT)
    - Paragraph packing with configurable length limits and overlap
    - Context breadcrumb preservation across split sections
    """

    SUPPORTED_TEXT_EXTENSIONS: Set[str] = {
        ".txt",
        ".md",
        ".markdown",
        ".text",
        ".memo",
    }
    SUPPORTED_DOCX_EXTENSIONS: Set[str] = {".docx"}

    # Default regex patterns for structural university boundaries
    DEFAULT_BOUNDARY_PATTERNS: List[Tuple[str, str]] = [
        # Markdown headers (# Header 1, ## Header 2, etc.)
        ("markdown_header", r"^(#{1,6})\s+(.+)$"),
        # Department / Faculty / Prodi administrative boundaries
        (
            "department_boundary",
            r"^(?:DEPARTMENT|PROGRAM\s+STUDI|PRODI|JURUSAN|FAKULTAS|FACULTY|DIVISI|BAGIAN|UNIT|INSTITUT|SEKOLAH)\s*[:\-–—]\s*(.+)$",
        ),
        # Administrative memo and formal notice boundaries
        (
            "administrative_heading",
            r"^(?:MEMORANDUM|MEMO|SURAT\s+EDARAN|SURAT\s+KEPUTUSAN|PERIHAL|SUBJECT)\s*[:\-–—]?\s*(.*)$",
        ),
        # Academic chapter / section / legal policy articles
        (
            "chapter_section",
            r"^(?:BAB\s+[0-9IVXLCDM]+|SECTION\s+[0-9IVXLCDM]+|PASAL\s+[0-9]+|LAMPIRAN\s+[0-9IVXLCDM]*)\b.*$",
        ),
        # Numbered major headings, e.g., '1. KETENTUAN UMUM' or 'II. PERSYARATAN JADWAL'
        (
            "numbered_heading",
            r"^([0-9IVXLCDM]+\.)\s+([A-Z0-9\s,\-–—]{3,})$",
        ),
    ]

    def __init__(
        self,
        max_chunk_chars: int = 2000,
        min_chunk_chars: int = 100,
        overlap_chars: int = 150,
        preserve_context_breadcrumbs: bool = True,
        custom_boundary_patterns: Optional[List[Tuple[str, str]]] = None,
    ) -> None:
        """
        Initialize the TextSlicer.

        Args:
            max_chunk_chars: Maximum character limit per chunk (approx. 500 tokens).
            min_chunk_chars: Minimum character threshold to prevent tiny fragmented chunks.
            overlap_chars: Character overlap when splitting long contiguous paragraphs.
            preserve_context_breadcrumbs: If True, prepends section title breadcrumb to subchunks.
            custom_boundary_patterns: Optional list of (boundary_type, regex_pattern) tuples.
        """
        if max_chunk_chars < 200:
            raise ValueError(f"max_chunk_chars must be >= 200, got: {max_chunk_chars}")
        if min_chunk_chars >= max_chunk_chars:
            raise ValueError("min_chunk_chars must be strictly less than max_chunk_chars")

        self.max_chunk_chars = max_chunk_chars
        self.min_chunk_chars = min_chunk_chars
        self.overlap_chars = max(0, overlap_chars)
        self.preserve_context_breadcrumbs = preserve_context_breadcrumbs

        # Compile boundary regex patterns
        patterns_to_compile = list(self.DEFAULT_BOUNDARY_PATTERNS)
        if custom_boundary_patterns:
            patterns_to_compile.extend(custom_boundary_patterns)

        self._compiled_boundaries: List[Tuple[str, Pattern[str]]] = [
            (b_type, re.compile(pat, re.IGNORECASE | re.MULTILINE))
            for b_type, pat in patterns_to_compile
        ]

    def slice_file(self, file_path: Union[str, Path]) -> List[TextChunk]:
        """
        Read and slice a text or docx document file.

        Args:
            file_path: Path to target text document.

        Returns:
            List of semantic TextChunk objects.
        """
        path = self._validate_file_path(file_path)
        ext = path.suffix.lower()

        if ext in self.SUPPORTED_DOCX_EXTENSIONS:
            content = self._read_docx(path)
        elif ext in self.SUPPORTED_TEXT_EXTENSIONS:
            content = path.read_text(encoding="utf-8", errors="replace")
        else:
            # Fallback: attempt to read as utf-8 plain text
            try:
                content = path.read_text(encoding="utf-8", errors="replace")
            except Exception as exc:
                raise ValueError(
                    f"Unsupported or unreadable file format '{ext}' for file {path}: {exc}"
                ) from exc

        return self.slice_text(content, source_file=path)

    def slice_text(
        self,
        text: str,
        source_file: Optional[Union[str, Path]] = None,
    ) -> List[TextChunk]:
        """
        Split a raw text string into semantic chunks along boundaries.

        Args:
            text: Raw input text content.
            source_file: Optional path to origin document for traceability.

        Returns:
            List of TextChunk objects.
        """
        resolved_source = Path(source_file).resolve() if source_file is not None else None
        clean_text = text.strip()
        if not clean_text:
            return []

        # 1. Segment text into initial sections via boundary patterns
        sections = self._partition_into_sections(clean_text)

        # 2. Refine sections: pack or split by paragraphs up to max_chunk_chars
        raw_chunks: List[Dict[str, Any]] = []
        for sec in sections:
            sec_chunks = self._chunk_section(sec)
            raw_chunks.extend(sec_chunks)

        if not raw_chunks:
            return []

        # 3. Assemble final TextChunk objects with total_chunks count
        total_chunks = len(raw_chunks)
        final_chunks: List[TextChunk] = []

        for idx, item in enumerate(raw_chunks):
            content_str = item["content"].strip()
            char_count = len(content_str)
            token_estimate = max(1, char_count // 4)

            metadata = dict(item.get("metadata", {}))
            metadata["subchunk_index"] = item.get("subchunk_index", 0)
            metadata["subchunk_total"] = item.get("subchunk_total", 1)

            final_chunks.append(
                TextChunk(
                    chunk_index=idx,
                    total_chunks=total_chunks,
                    title=item.get("title"),
                    content=content_str,
                    char_count=char_count,
                    token_estimate=token_estimate,
                    source_file=resolved_source,
                    metadata=metadata,
                )
            )

        return final_chunks

    def _validate_file_path(self, file_path: Union[str, Path]) -> Path:
        """Validate existence and non-zero size."""
        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")
        if not path.is_file():
            raise ValueError(f"Path is not a regular file: {path}")
        if path.stat().st_size == 0:
            raise ValueError(f"File is empty (0 bytes): {path}")
        return path

    def _read_docx(self, file_path: Path) -> str:
        """Extract text paragraphs and tables from a .docx file."""
        if docx is None:
            raise ImportError("python-docx is required to parse .docx documents.")

        doc = docx.Document(file_path)
        lines: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                # If paragraph style is a Heading, preserve as markdown header
                if p.style and p.style.name.startswith("Heading "):
                    try:
                        level = int(p.style.name.split()[-1])
                        lines.append(f"{'#' * min(level, 6)} {text}")
                    except ValueError:
                        lines.append(f"## {text}")
                else:
                    lines.append(text)

        # Also extract text from tables inside docx
        for table in doc.tables:
            table_lines: List[str] = []
            for row in table.rows:
                row_vals = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(row_vals):
                    table_lines.append(" | ".join(row_vals))
            if table_lines:
                lines.append("\n".join(table_lines))

        return "\n\n".join(lines)

    def _partition_into_sections(self, text: str) -> List[Dict[str, Any]]:
        """
        Scan text line-by-line to identify semantic boundaries and group into sections.
        """
        lines = text.splitlines()
        sections: List[Dict[str, Any]] = []

        current_title: Optional[str] = None
        current_boundary_type: Optional[str] = None
        current_lines: List[str] = []

        for line in lines:
            trimmed = line.strip()
            detected_boundary = self._match_boundary(trimmed)

            if detected_boundary is not None:
                b_type, b_title = detected_boundary
                # Flush previous section if it has content
                if current_lines:
                    sec_body = "\n".join(current_lines).strip()
                    if sec_body:
                        sections.append(
                            {
                                "title": current_title,
                                "boundary_type": current_boundary_type,
                                "text": sec_body,
                            }
                        )
                    current_lines = []

                current_title = b_title
                current_boundary_type = b_type
                current_lines.append(line)
            else:
                current_lines.append(line)

        # Flush final accumulated section
        if current_lines:
            sec_body = "\n".join(current_lines).strip()
            if sec_body:
                sections.append(
                    {
                        "title": current_title,
                        "boundary_type": current_boundary_type,
                        "text": sec_body,
                    }
                )

        # Guard against zero sections
        if not sections and text:
            sections.append(
                {
                    "title": None,
                    "boundary_type": "plain_text",
                    "text": text,
                }
            )

        return sections

    def _match_boundary(self, line: str) -> Optional[Tuple[str, str]]:
        """Check if a line matches any registered boundary regex pattern."""
        if not line:
            return None

        for b_type, pattern in self._compiled_boundaries:
            match = pattern.match(line)
            if match:
                # Extract descriptive title from match groups
                groups = match.groups()
                if len(groups) >= 2:
                    title = groups[1].strip()
                elif len(groups) == 1:
                    title = groups[0].strip()
                else:
                    title = line.strip()
                return b_type, title

        return None

    def _chunk_section(self, section: Dict[str, Any]) -> List[Dict[str, Any]]:
        """
        Subdivide a semantic section by double line breaks (paragraphs) or length limit.
        """
        sec_text = section["text"]
        title = section["title"]
        b_type = section["boundary_type"]

        # If section already fits comfortably within max_chunk_chars, return intact
        if len(sec_text) <= self.max_chunk_chars:
            return [
                {
                    "title": title,
                    "content": sec_text,
                    "subchunk_index": 0,
                    "subchunk_total": 1,
                    "metadata": {
                        "boundary_type": b_type,
                        "section_title": title,
                    },
                }
            ]

        # Section exceeds max length: split into paragraphs by double line breaks
        raw_paragraphs = [p.strip() for p in re.split(r"\n\s*\n", sec_text) if p.strip()]
        if not raw_paragraphs:
            raw_paragraphs = [sec_text]

        # Pack paragraphs into batches up to max_chunk_chars
        packed_chunks: List[str] = []
        curr_paras: List[str] = []
        curr_len = 0

        for para in raw_paragraphs:
            para_len = len(para)

            # If a single paragraph is larger than max_chunk_chars, split by sentence or window
            if para_len > self.max_chunk_chars:
                if curr_paras:
                    packed_chunks.append("\n\n".join(curr_paras))
                    curr_paras = []
                    curr_len = 0

                sub_parts = self._split_long_paragraph(para)
                packed_chunks.extend(sub_parts)
                continue

            if curr_len + para_len + 2 > self.max_chunk_chars:
                if curr_paras:
                    packed_chunks.append("\n\n".join(curr_paras))
                    curr_paras = [para]
                    curr_len = para_len
                else:
                    curr_paras = [para]
                    curr_len = para_len
            else:
                curr_paras.append(para)
                curr_len += para_len + 2

        if curr_paras:
            packed_chunks.append("\n\n".join(curr_paras))

        # Consolidate trailing subchunk if smaller than min_chunk_chars and fits in previous chunk
        if len(packed_chunks) > 1 and len(packed_chunks[-1]) < self.min_chunk_chars:
            if len(packed_chunks[-2]) + len(packed_chunks[-1]) + 2 <= self.max_chunk_chars:
                packed_chunks[-2] = f"{packed_chunks[-2]}\n\n{packed_chunks[-1]}"
                packed_chunks.pop()

        # Build subchunk metadata with context preservation
        subchunk_total = len(packed_chunks)
        results: List[Dict[str, Any]] = []

        for sub_idx, chunk_text in enumerate(packed_chunks):
            # Prepend breadcrumb if split and enabled
            final_content = chunk_text
            if self.preserve_context_breadcrumbs and title and subchunk_total > 1:
                if sub_idx > 0 and not chunk_text.startswith(f"[{title}]"):
                    final_content = f"[{title} (Lanjutan {sub_idx+1}/{subchunk_total})]\n\n{chunk_text}"

            results.append(
                {
                    "title": title,
                    "content": final_content,
                    "subchunk_index": sub_idx,
                    "subchunk_total": subchunk_total,
                    "metadata": {
                        "boundary_type": b_type,
                        "section_title": title,
                        "is_subdivided": subchunk_total > 1,
                    },
                }
            )

        return results

    def _split_long_paragraph(self, para: str) -> List[str]:
        """Split a gigantic paragraph using sentence boundaries or sliding window."""
        # Split by sentence terminators (.!?)
        sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", para) if s.strip()]
        if len(sentences) <= 1:
            return self._sliding_window_split(para)

        # Pack sentences
        sub_chunks: List[str] = []
        curr_sentences: List[str] = []
        curr_len = 0

        for s in sentences:
            s_len = len(s)
            # If an individual sentence exceeds max_chunk_chars, split with sliding window
            if s_len > self.max_chunk_chars:
                if curr_sentences:
                    sub_chunks.append(" ".join(curr_sentences))
                    curr_sentences = []
                    curr_len = 0
                sub_chunks.extend(self._sliding_window_split(s))
                continue

            if curr_len + s_len + 1 > self.max_chunk_chars and curr_sentences:
                sub_chunks.append(" ".join(curr_sentences))
                curr_sentences = [s]
                curr_len = s_len
            else:
                curr_sentences.append(s)
                curr_len += s_len + 1

        if curr_sentences:
            sub_chunks.append(" ".join(curr_sentences))

        return sub_chunks

    def _sliding_window_split(self, text: str) -> List[str]:
        """Fallback character chunking with overlap for text exceeding max_chunk_chars."""
        parts: List[str] = []
        start = 0
        step = max(100, self.max_chunk_chars - self.overlap_chars)
        while start < len(text):
            end = min(start + self.max_chunk_chars, len(text))
            parts.append(text[start:end])
            if end >= len(text):
                break
            start += step
        return parts
