"""Tests for pipeline audit fixes across UniTime AI Ingestion Gateway.

Verifies:
1. Excel table header banner detection and academic keyword scoring.
2. Multi-sheet Excel monotonic chunk IDs, sheet_name metadata, and markdown prefixes.
3. Scanned PDF fallback image rendering at 200 DPI.
4. Document name and page context inclusion in multimodal vision user prompts.
5. Secondary fallback to extracted_text metadata when vision extraction fails.
6. Global instructor teaching share normalization and single-lead enforcement.
"""

from pathlib import Path
from typing import Any, Dict
from unittest.mock import MagicMock, patch

import openpyxl
import pymupdf
import pytest

from agent.nodes import extract_chunk_node
from agent.state import create_initial_state
from core.extractor import (
    AnthropicExtractor,
    BaseExtractor,
    ExtractionResult,
    GeminiExtractor,
    OpenAIExtractor,
    _build_vision_user_prompt,
)
from core.merger import Merger
from core.slicer import DocumentSlicer
from slicers.table_slicer import TableSlicer


# =============================================================================
# 1. Excel Header Banner Detection & Academic Keyword Scoring
# =============================================================================


def test_excel_header_banner_detection_on_sample_file() -> None:
    """Verify sample_inputs/jadwal_kuliah_if.xlsx detects row 4 headers instead of row 1 banner."""
    sample_file = Path(__file__).resolve().parent.parent / "sample_inputs" / "jadwal_kuliah_if.xlsx"
    assert sample_file.exists(), f"Sample file not found: {sample_file}"

    slicer = TableSlicer()
    chunks = list(slicer.slice(sample_file))
    assert len(chunks) >= 1

    first_chunk = chunks[0]
    expected_headers = [
        "Kode MK",
        "Mata Kuliah",
        "SKS",
        "Komponen",
        "Kelas",
        "Kapasitas",
        "Hari",
        "Waktu",
        "Gedung",
        "Ruang",
        "Dosen Pengampu",
    ]
    assert first_chunk.headers == expected_headers
    # Ensure row 1 banner text is NOT treated as a column header
    for h in first_chunk.headers:
        assert "JADWAL PERKULIAHAN" not in h
        assert "INSTITUT TEKNOLOGI" not in h

    # Verify first data row corresponds to actual course data
    assert first_chunk.rows[0][0] == "IF2110"
    assert first_chunk.rows[0][1] == "Algoritma & Struktur Data"
    assert first_chunk.metadata.get("sheet_name") == "Jadwal Ganjil 2024"


def test_excel_header_banner_detection_synthetic_banners(tmp_path: Path) -> None:
    """Verify banner detection rejects multiple uniform rows and prioritizes academic keywords."""
    wb_path = tmp_path / "banner_test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Schedule"

    # Row 1: Merged title banner across 8 columns
    ws.append(["UNIVERSITAS INDONESIA - FAKULTAS ILMU KOMPUTER"] * 8)
    # Row 2: Merged subtitle banner across 8 columns
    ws.append(["SEMESTER GENAP 2024/2025"] * 8)
    # Row 3: Actual academic table headers
    ws.append(["Kode", "Mata Kuliah", "SKS", "Kelas", "Hari", "Waktu", "Ruang", "Dosen"])
    # Row 4..5: Data rows
    ws.append(["CS101", "Dasar Pemrograman", 3, "A", "Senin", "08.00-10.00", "Lab-1", "Dosen A"])
    ws.append(["CS102", "Struktur Data", 3, "B", "Selasa", "10.00-12.00", "Lab-2", "Dosen B"])

    wb.save(wb_path)
    wb.close()

    slicer = TableSlicer()
    chunks = slicer.slice_all(wb_path)
    assert len(chunks) == 1
    assert chunks[0].headers == [
        "Kode",
        "Mata Kuliah",
        "SKS",
        "Kelas",
        "Hari",
        "Waktu",
        "Ruang",
        "Dosen",
    ]
    assert len(chunks[0].rows) == 2
    assert chunks[0].rows[0][0] == "CS101"


# =============================================================================
# 2. Multi-Sheet Monotonic Chunk IDs & Sheet Context
# =============================================================================


def test_document_slicer_multi_sheet_monotonic_chunk_ids(tmp_path: Path) -> None:
    """Ensure chunk_id is globally monotonic and sheet_name is preserved in metadata and prefix."""
    wb_path = tmp_path / "multi_prodi.xlsx"
    wb = openpyxl.Workbook()

    sheets = ["Prodi_IF", "Prodi_SI", "Prodi_EL"]
    for i, s_name in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet(title=s_name)
        ws.title = s_name
        ws.append(["Kode MK", "Mata Kuliah", "SKS", "Dosen"])
        ws.append([f"CS{i}01", f"Course A{i}", 3, f"Lecturer A{i}"])
        ws.append([f"CS{i}02", f"Course B{i}", 3, f"Lecturer B{i}"])

    wb.save(wb_path)
    wb.close()

    # Slicing with excel_rows_per_chunk=1 yields 2 chunks per sheet = 6 chunks total
    slicer = DocumentSlicer(excel_rows_per_chunk=1)
    chunks = slicer.slice_file(wb_path)

    assert len(chunks) == 6
    for idx, c in enumerate(chunks, start=1):
        assert c.chunk_id == idx
        assert c.total_chunks == 6
        assert "sheet_name" in c.metadata
        sheet = c.metadata["sheet_name"]
        assert c.text_content is not None
        assert c.text_content.startswith(f"# Sheet: {sheet}\n\n")

    # Verify sheet names match expectations in order
    assert [c.metadata["sheet_name"] for c in chunks] == [
        "Prodi_IF",
        "Prodi_IF",
        "Prodi_SI",
        "Prodi_SI",
        "Prodi_EL",
        "Prodi_EL",
    ]


# =============================================================================
# 3. Scanned PDF Fallback DPI
# =============================================================================


def test_document_slicer_scanned_pdf_fallback_dpi(tmp_path: Path) -> None:
    """Verify pages with no extractable text render image at configured pdf_dpi (200 DPI)."""
    pdf_path = tmp_path / "scanned_doc.pdf"
    doc = pymupdf.open()
    # Create empty page (simulating scanned non-searchable PDF page)
    page = doc.new_page(width=595, height=842)  # A4 at 72 pt/inch
    doc.save(str(pdf_path))
    doc.close()

    slicer = DocumentSlicer(render_pdf_as_images=False, pdf_dpi=200)
    chunks = slicer.slice_file(pdf_path)

    assert len(chunks) == 1
    chunk = chunks[0]
    assert chunk.chunk_type == "image"
    assert chunk.metadata.get("note") == "fallback_ocr"
    assert chunk.image_bytes is not None

    # Load pixmap from rendered bytes to verify high-res 200 DPI rendering
    rendered_pix = pymupdf.Pixmap(chunk.image_bytes)
    # At 72 DPI, A4 is 595x842. At 200 DPI, width is round(595 * 200 / 72) = 1653
    assert rendered_pix.width > 1500
    assert rendered_pix.height > 2000
    assert rendered_pix.alpha == 0


# =============================================================================
# 4. Multimodal Vision Prompt Context
# =============================================================================


def test_build_vision_user_prompt_formatting() -> None:
    """Verify _build_vision_user_prompt prefixes source document and page context."""
    base = "Extract academic schedule into JSON."

    # Context with document name and page
    prompt_with_page = _build_vision_user_prompt(
        base,
        {"sourceDocumentName": "jadwal.pdf", "page": 2, "total_pages": 5},
    )
    assert "Source Document: jadwal.pdf" in prompt_with_page
    assert "Page: 2/5" in prompt_with_page
    assert prompt_with_page.endswith(f"\n\n{base}")

    # Context with document name and chunk_id
    prompt_with_chunk = _build_vision_user_prompt(
        base,
        {"sourceDocumentName": "schedule.xlsx", "chunk_id": 1, "total_chunks": 3},
    )
    assert "Source Document: schedule.xlsx" in prompt_with_chunk
    assert "Chunk/Page: 1/3" in prompt_with_chunk

    # Empty context returns base prompt unchanged
    assert _build_vision_user_prompt(base, None) == base
    assert _build_vision_user_prompt(base, {}) == base


def test_extractors_incorporate_document_context_in_vision_prompts() -> None:
    """Verify Gemini, OpenAI, and Anthropic vision requests include document metadata in prompt."""
    dummy_bytes = b"fake_image_bytes"
    context = {"sourceDocumentName": "kurikulum_2024.pdf", "page": 3, "total_pages": 10}

    # 1. Gemini
    gemini = GeminiExtractor(api_key="test-gemini-key")
    with patch.object(gemini, "_send_request") as mock_gemini_send:
        mock_gemini_send.return_value = ExtractionResult(
            payload={}, provider="gemini", model="gemini", raw_response="{}"
        )
        gemini.extract_image(dummy_bytes, context=context)
        sent_body = mock_gemini_send.call_args[0][0]
        prompt_text = sent_body["contents"][0]["parts"][0]["text"]
        assert "Source Document: kurikulum_2024.pdf" in prompt_text
        assert "Page: 3/10" in prompt_text

    # 2. OpenAI
    openai_ext = OpenAIExtractor(api_key="test-openai-key")
    with patch.object(openai_ext, "_send_request") as mock_openai_send:
        mock_openai_send.return_value = ExtractionResult(
            payload={}, provider="openai", model="gpt-4o", raw_response="{}"
        )
        openai_ext.extract_image(dummy_bytes, context=context)
        sent_body = mock_openai_send.call_args[0][0]
        user_msg = sent_body["messages"][1]["content"]
        prompt_text = user_msg[0]["text"]
        assert "Source Document: kurikulum_2024.pdf" in prompt_text
        assert "Page: 3/10" in prompt_text

    # 3. Anthropic
    anthropic_ext = AnthropicExtractor(api_key="test-anthropic-key")
    with patch.object(anthropic_ext, "_send_request") as mock_anthropic_send:
        mock_anthropic_send.return_value = ExtractionResult(
            payload={}, provider="anthropic", model="claude", raw_response="{}"
        )
        anthropic_ext.extract_image(dummy_bytes, context=context)
        sent_body = mock_anthropic_send.call_args[0][0]
        user_msg = sent_body["messages"][0]["content"]
        prompt_text = user_msg[1]["text"]
        assert "Source Document: kurikulum_2024.pdf" in prompt_text
        assert "Page: 3/10" in prompt_text


# =============================================================================
# 5. Secondary Fallback on Vision Extraction Failure
# =============================================================================


def test_extract_chunk_node_secondary_text_fallback_on_vision_error() -> None:
    """Verify extract_chunk_node falls back to extracted_text metadata when vision extraction fails."""
    slices = [
        {
            "chunk_id": 1,
            "total_chunks": 1,
            "chunk_type": "image",
            "source_filename": "scanned_memo.pdf",
            "image_bytes": b"mock_corrupt_or_rate_limited_image",
            "mime_type": "image/png",
            "metadata": {
                "page": 1,
                "extracted_text": "Mata Kuliah: IF2110 Algoritma, Kelas: K01, Dosen: Dr. Ayu",
            },
        }
    ]

    state = create_initial_state(document_path="scanned_memo.pdf")
    state["slices"] = slices
    state["memory_context"] = {"provider": "mock"}

    mock_extractor = MagicMock(spec=BaseExtractor)
    mock_extractor.model = "mock-vision-test"
    # Vision extraction raises an error
    mock_extractor.extract_image.side_effect = RuntimeError("Vision API 503 Overloaded")
    # Text fallback succeeds
    fallback_payload = {
        "courses": [
            {
                "courseNumber": "IF2110",
                "configurations": [
                    {
                        "name": "Default",
                        "subparts": [
                            {
                                "type": "Lecture",
                                "classes": [
                                    {
                                        "sectionName": "K01",
                                        "instructors": [{"id": "1", "name": "Dr. Ayu", "isLead": True}],
                                    }
                                ],
                            }
                        ],
                    }
                ],
            }
        ]
    }
    mock_extractor.extract_text.return_value = ExtractionResult(
        payload=fallback_payload,
        provider="mock",
        model="mock-text-fallback",
        raw_response="{}",
    )

    with patch("agent.nodes.get_extractor", return_value=mock_extractor):
        result = extract_chunk_node(state)

    # Extraction must have succeeded via secondary text fallback
    assert len(result.get("partial_payloads", [])) == 1
    assert result["partial_payloads"][0] == fallback_payload
    mock_extractor.extract_image.assert_called_once()
    mock_extractor.extract_text.assert_called_once()
    # Check that text passed to fallback matches metadata extracted_text
    called_text = mock_extractor.extract_text.call_args[0][0]
    assert "IF2110 Algoritma" in called_text


def test_extract_chunk_node_vision_error_without_fallback_text() -> None:
    """Verify extract_chunk_node logs ERROR when vision extraction fails and no extracted_text is present."""
    slices = [
        {
            "chunk_id": 1,
            "total_chunks": 1,
            "chunk_type": "image",
            "source_filename": "photo.png",
            "image_bytes": b"corrupted_bytes",
            "mime_type": "image/png",
            "metadata": {"page": 1},  # No extracted_text
        }
    ]

    state = create_initial_state(document_path="photo.png")
    state["slices"] = slices
    state["memory_context"] = {"provider": "mock"}

    mock_extractor = MagicMock(spec=BaseExtractor)
    mock_extractor.model = "mock-vision-test"
    mock_extractor.extract_image.side_effect = RuntimeError("Vision API unrecoverable error")

    with patch("agent.nodes.get_extractor", return_value=mock_extractor):
        result = extract_chunk_node(state)

    # No payload extracted, audit log records ERROR
    assert len(result.get("partial_payloads", [])) == 0
    assert any(log["level"] == "ERROR" for log in result.get("audit_logs", []))


# =============================================================================
# 6. Global Instructor Load & Lead Normalization
# =============================================================================


def test_merger_global_instructor_share_and_lead_normalization() -> None:
    """Verify _standardize_global_instructors normalizes multi-instructor shares and lead flag."""
    payload: Dict[str, Any] = {
        "courses": [
            {
                "courseNumber": "IF2110",
                "configurations": [
                    {
                        "name": "Default",
                        "subparts": [
                            {
                                "type": "Lecture",
                                "classes": [
                                    {
                                        "sectionName": "K01",
                                        # 3 instructors without shares, multiple marked lead
                                        "instructors": [
                                            {"id": "INS01", "name": "Prof A", "isLead": True},
                                            {"id": "INS02", "name": "Dr B", "isLead": True},
                                            {"id": "INS03", "name": "Dr C", "isLead": False},
                                        ],
                                    },
                                    {
                                        "sectionName": "K02",
                                        # 2 instructors with no lead specified
                                        "instructors": [
                                            {"id": "INS04", "name": "Dr D"},
                                            {"id": "INS05", "name": "Dr E"},
                                        ],
                                    },
                                    {
                                        "sectionName": "K03",
                                        # Single instructor
                                        "instructors": [
                                            {"id": "INS06", "name": "Prof Single"},
                                        ],
                                    },
                                    {
                                        "sectionName": "K04",
                                        # Pre-existing valid 60/40 shares, but both marked isLead: True
                                        "instructors": [
                                            {"id": "INS07", "name": "Lead X", "sharePercentage": 60, "isLead": True},
                                            {"id": "INS08", "name": "Member Y", "sharePercentage": 40, "isLead": True},
                                        ],
                                    },
                                ],
                            }
                        ],
                    }
                ],
            }
        ]
    }

    merger = Merger()
    merged = merger.merge([payload])

    classes = merged["courses"][0]["configurations"][0]["subparts"][0]["classes"]

    # Class K01: N=3 -> lead=34, member=33, exactly 1 lead
    k01_ins = classes[0]["instructors"]
    assert sum(i["sharePercentage"] for i in k01_ins) == 100
    assert sum(1 for i in k01_ins if i["isLead"] is True) == 1
    assert k01_ins[0]["isLead"] is True
    assert k01_ins[1]["isLead"] is False
    assert k01_ins[2]["isLead"] is False
    assert k01_ins[0]["sharePercentage"] == 34
    assert k01_ins[1]["sharePercentage"] == 33
    assert k01_ins[2]["sharePercentage"] == 33

    # Class K02: N=2 -> lead=50, member=50, first designated as lead
    k02_ins = classes[1]["instructors"]
    assert sum(i["sharePercentage"] for i in k02_ins) == 100
    assert sum(1 for i in k02_ins if i["isLead"] is True) == 1
    assert k02_ins[0]["isLead"] is True
    assert k02_ins[1]["isLead"] is False
    assert k02_ins[0]["sharePercentage"] == 50
    assert k02_ins[1]["sharePercentage"] == 50

    # Class K03: N=1 -> lead=100, isLead=True
    k03_ins = classes[2]["instructors"]
    assert len(k03_ins) == 1
    assert k03_ins[0]["isLead"] is True
    assert k03_ins[0]["sharePercentage"] == 100

    # Class K04: Preserves 60/40 valid shares, but demotes second lead
    k04_ins = classes[3]["instructors"]
    assert k04_ins[0]["sharePercentage"] == 60
    assert k04_ins[1]["sharePercentage"] == 40
    assert sum(1 for i in k04_ins if i["isLead"] is True) == 1
    assert k04_ins[0]["isLead"] is True
    assert k04_ins[1]["isLead"] is False
